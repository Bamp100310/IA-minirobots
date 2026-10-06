"""Extrae el texto de un documento conservando la página de origen.

Uso:  python extraer_texto.py <documento.(pdf|docx|txt|md)>

La salida es texto plano. Cada página empieza con una línea '### pagina N';
los documentos sin páginas (txt, md, docx) se reportan como página 1.
"""
import shutil
import subprocess
import sys
from pathlib import Path


def paginas_pdf(ruta):
    try:
        from pypdf import PdfReader
        return [p.extract_text() or "" for p in PdfReader(str(ruta)).pages]
    except ImportError:
        if shutil.which("pdftotext") is None:
            sys.exit("Se necesita 'pypdf' (pip install pypdf) o 'pdftotext' para leer PDF.")
        texto = subprocess.run(["pdftotext", "-layout", str(ruta), "-"],
                               capture_output=True, text=True, check=True).stdout
        return texto.split("\f")


def paginas_docx(ruta):
    try:
        import docx
    except ImportError:
        sys.exit("Se necesita 'python-docx' (pip install python-docx) para leer DOCX.")
    documento = docx.Document(str(ruta))
    lineas = [p.text for p in documento.paragraphs]
    for tabla in documento.tables:
        for fila in tabla.rows:
            lineas.append(" | ".join(c.text for c in fila.cells))
    return ["\n".join(lineas)]


def extraer(ruta):
    ruta = Path(ruta)
    sufijo = ruta.suffix.lower()
    if sufijo == ".pdf":
        paginas = paginas_pdf(ruta)
    elif sufijo == ".docx":
        paginas = paginas_docx(ruta)
    elif sufijo in (".txt", ".md", ".markdown"):
        paginas = [ruta.read_text(encoding="utf-8", errors="replace")]
    else:
        sys.exit(f"Formato no soportado: {sufijo}")
    return "\n".join(f"### pagina {i}\n{texto.rstrip()}" for i, texto in enumerate(paginas, 1))


if __name__ == "__main__":
    if len(sys.argv) != 2:
        sys.exit(__doc__)
    sys.stdout.reconfigure(encoding="utf-8")
    print(extraer(sys.argv[1]))
